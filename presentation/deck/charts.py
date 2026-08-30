"""Native PowerPoint charts, styled to the template.

Charts are emitted as real chart parts (not images) so the numbers stay
editable in PowerPoint and stay crisp when projected.
"""

from __future__ import annotations

from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_TICK_LABEL_POSITION
from pptx.util import Inches, Pt

from .theme import Color, Font, rgb


def add_bar_chart(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    categories: list[str],
    values: list[float],
    series_name: str = "Nilai",
    number_format: str = '0.00"%"',
    horizontal: bool = True,
    highlight: int | None = None,
    labels: list[str] | None = None,
    gap_width: int = 55,
):
    """A single-series bar chart in navy/gold.

    `highlight` is the index of the bar to paint navy instead of gold —
    used to mark the winning model. `labels` overrides the numeric data
    labels with literal text, which is how the Indonesian decimal comma is
    guaranteed regardless of the viewer's locale (ITB Pedoman §VIII.3).
    """
    data = CategoryChartData()
    data.categories = categories
    data.add_series(series_name, values)

    kind = XL_CHART_TYPE.BAR_CLUSTERED if horizontal else XL_CHART_TYPE.COLUMN_CLUSTERED
    frame = slide.shapes.add_chart(kind, Inches(x), Inches(y), Inches(w), Inches(h), data)
    chart = frame.chart
    chart.has_title = False
    chart.has_legend = False

    plot = chart.plots[0]
    plot.gap_width = gap_width
    plot.overlap = 0
    plot.vary_by_categories = False

    series = plot.series[0]
    series.format.fill.solid()
    series.format.fill.fore_color.rgb = rgb(Color.GOLD)
    series.format.line.fill.background()

    if highlight is not None and 0 <= highlight < len(values):
        point = series.points[highlight]
        point.format.fill.solid()
        point.format.fill.fore_color.rgb = rgb(Color.NAVY)
        point.format.line.fill.background()

    plot.has_data_labels = True
    dlbls = plot.data_labels
    dlbls.show_value = True
    dlbls.number_format = number_format
    dlbls.number_format_is_linked = False
    dlbls.position = XL_LABEL_POSITION.OUTSIDE_END
    dlbls.font.size = Pt(11)
    dlbls.font.bold = True
    dlbls.font.name = Font.BODY
    dlbls.font.color.rgb = rgb(Color.NAVY)

    if labels:
        for idx, text in enumerate(labels[: len(values)]):
            label = series.points[idx].data_label
            label.position = XL_LABEL_POSITION.OUTSIDE_END
            frame = label.text_frame
            frame.text = text
            run = frame.paragraphs[0].runs[0]
            run.font.size = Pt(11)
            run.font.bold = True
            run.font.name = Font.BODY
            run.font.color.rgb = rgb(Color.NAVY)

    cat_axis = chart.category_axis
    cat_axis.has_major_gridlines = False
    cat_axis.tick_labels.font.size = Pt(11)
    cat_axis.tick_labels.font.name = Font.BODY
    cat_axis.tick_labels.font.color.rgb = rgb(Color.INK)
    cat_axis.format.line.color.rgb = rgb(Color.MUTED)

    # The value axis is suppressed: every bar already carries its own label,
    # and an axis of raw floats would print machine-rounded tick values.
    val_axis = chart.value_axis
    val_axis.visible = False
    val_axis.has_major_gridlines = False
    val_axis.tick_labels.position = XL_TICK_LABEL_POSITION.NONE
    val_axis.format.line.fill.background()

    lo, hi = min(values), max(values)
    span = (hi - lo) or max(abs(hi), 1.0)
    val_axis.minimum_scale = max(0.0, lo - span * 0.45)
    val_axis.maximum_scale = hi + span * 0.18

    return chart
