"""Slide-level layouts: cover, content, section divider, closing."""

from __future__ import annotations

from pptx.util import Inches

from . import shapes
from .layout import render_stack
from .theme import Color, Font, Grid, Size, Stroke

G = Grid


# --------------------------------------------------------------------------
# Shared chrome
# --------------------------------------------------------------------------
def add_footer(slide, text: str, page: int | None) -> None:
    shapes.label(
        slide, G.MARGIN, G.FOOTER_Y, G.CONTENT_W - 0.83, G.FOOTER_H, text,
        size=Size.FOOTER, color=Color.MUTED,
    )
    if page is not None:
        shapes.label(
            slide, G.PAGENO_X, G.FOOTER_Y, G.PAGENO_W, G.FOOTER_H, str(page),
            size=Size.FOOTER, color=Color.MUTED, align="right",
        )


def add_header(slide, eyebrow: str | None, title: str) -> None:
    if eyebrow:
        shapes.label(
            slide, G.MARGIN, G.EYEBROW_Y, G.CONTENT_W, G.EYEBROW_H, eyebrow.upper(),
            size=Size.EYEBROW, bold=True, color=Color.GOLD,
        )
    shapes.label(
        slide, G.MARGIN, G.TITLE_Y, G.CONTENT_W, G.TITLE_H, title,
        size=Size.TITLE, font=Font.DISPLAY, bold=True, color=Color.NAVY, accent=Color.GOLD,
    )


def blank(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    for shape in list(slide.shapes):
        shape._element.getparent().remove(shape._element)
    return slide


# --------------------------------------------------------------------------
# Cover
# --------------------------------------------------------------------------
def cover(prs, spec: dict, meta: dict) -> None:
    slide = blank(prs)
    shapes.set_background(slide, Color.NAVY)
    shapes.rect(slide, 0.28, 0.28, G.SLIDE_W - 0.56, G.SLIDE_H - 0.56,
                fill=Color.NAVY, line=Color.GOLD, line_w=Stroke.THICK)

    logo_h = spec.get("logo_h", 1.00)
    if spec.get("logo"):
        slide.shapes.add_picture(
            spec["logo"], Inches((G.SLIDE_W - logo_h) / 2), Inches(0.62), height=Inches(logo_h)
        )
    elif spec.get("badge"):
        bx, by, bs = (G.SLIDE_W - 1.00) / 2, 0.62, 1.00
        lines = spec["badge"] if isinstance(spec["badge"], list) else [spec["badge"]]
        shapes.rect(slide, bx, by, bs, bs, fill=Color.NAVY, line=Color.GOLD, line_w=Stroke.THIN)
        shapes.label(slide, bx, by, bs, bs, lines,
                     size=10, bold=True, align="center", anchor="middle", color=Color.GOLD)

    shapes.label(slide, 0.60, 1.78, G.SLIDE_W - 1.20, 0.34, meta["institution"].upper(),
                 size=16, font=Font.DISPLAY, bold=True, color=Color.WHITE, align="center")
    shapes.label(slide, 0.60, 2.12, G.SLIDE_W - 1.20, 0.30, meta["faculty"],
                 size=11.5, color=Color.ON_DARK_4, align="center")
    shapes.hrule(slide, (G.SLIDE_W - 2.40) / 2, 2.62, 2.40, color=Color.GOLD, width=Stroke.THIN)

    shapes.label(slide, 0.80, 2.85, G.SLIDE_W - 1.60, 1.20, spec["title"],
                 size=24, font=Font.DISPLAY, bold=True, color=Color.WHITE, align="center")
    shapes.label(slide, 0.80, 4.05, G.SLIDE_W - 1.60, 0.40, spec["subtitle"],
                 size=14, font=Font.DISPLAY, italic=True, color=Color.GOLD_LIGHT, align="center")
    shapes.label(slide, 0.80, 4.55, G.SLIDE_W - 1.60, 0.30, spec["kind"],
                 size=11, color=Color.ON_DARK_4, align="center")

    shapes.label(
        slide, 0.80, 4.98, G.SLIDE_W - 1.60, 0.36,
        f"**{meta['author']}**     ~~NIM {meta['nim']}~~",
        size=16, font=Font.DISPLAY, color=Color.WHITE, accent=Color.WHITE,
        muted=Color.GOLD_LIGHT, align="center",
    )
    shapes.label(
        slide, 0.80, 5.46, G.SLIDE_W - 1.60, 0.28,
        f"**Promotor   **~~{meta['promotor']}~~",
        size=11, color=Color.ON_DARK_2, accent=Color.GOLD, muted=Color.ON_DARK_2, align="center",
    )
    shapes.label(
        slide, 0.80, 5.74, G.SLIDE_W - 1.60, 0.28,
        "**Ko-promotor   **~~" + "   ·   ".join(meta["kopromotor"]) + "~~",
        size=11, color=Color.ON_DARK_2, accent=Color.GOLD, muted=Color.ON_DARK_2, align="center",
    )
    shapes.label(slide, 0.80, 6.42, G.SLIDE_W - 1.60, 0.30, spec["occasion"],
                 size=11, color=Color.ON_DARK_5, align="center")


# --------------------------------------------------------------------------
# Section divider
# --------------------------------------------------------------------------
def section(prs, spec: dict, meta: dict, page: int) -> None:
    slide = blank(prs)
    shapes.set_background(slide, Color.NAVY)
    shapes.label(slide, G.MARGIN + 0.45, 2.55, G.CONTENT_W - 0.90, 0.40,
                 spec.get("eyebrow", "").upper(), size=12, bold=True, color=Color.GOLD)
    shapes.hrule(slide, G.MARGIN + 0.45, 3.02, 2.40, color=Color.GOLD)
    shapes.label(slide, G.MARGIN + 0.45, 3.24, G.CONTENT_W - 0.90, 0.90, spec["title"],
                 size=38, font=Font.DISPLAY, bold=True, color=Color.WHITE)
    if spec.get("body"):
        shapes.label(slide, G.MARGIN + 0.45, 4.30, 8.40, 0.80, spec["body"],
                     size=13, color=Color.ON_DARK_2, accent=Color.WHITE)
    add_footer(slide, meta["footer"], page)


# --------------------------------------------------------------------------
# Content
# --------------------------------------------------------------------------
def content(prs, spec: dict, meta: dict, page: int) -> None:
    slide = blank(prs)
    shapes.set_background(slide, Color.WHITE)
    add_header(slide, spec.get("eyebrow"), spec["title"])
    top = spec.get("top", G.BAND_TOP)
    bottom = spec.get("bottom", G.BAND_BOTTOM)
    render_stack(slide, spec.get("blocks", []), G.MARGIN, top, G.CONTENT_W, bottom)
    add_footer(slide, meta["footer"], page)


# --------------------------------------------------------------------------
# Closing
# --------------------------------------------------------------------------
def closing(prs, spec: dict, meta: dict) -> None:
    slide = blank(prs)
    shapes.set_background(slide, Color.NAVY)
    shapes.rect(slide, 0.28, 0.28, G.SLIDE_W - 0.56, G.SLIDE_H - 0.56,
                fill=Color.NAVY, line=Color.GOLD, line_w=Stroke.THICK)
    shapes.label(slide, 0.60, 2.05, G.SLIDE_W - 1.20, 0.40, meta["institution"].upper(),
                 size=16, font=Font.DISPLAY, bold=True, color=Color.ON_DARK_4, align="center")
    shapes.hrule(slide, (G.SLIDE_W - 2.40) / 2, 2.66, 2.40, color=Color.GOLD)
    shapes.label(slide, 0.60, 2.95, G.SLIDE_W - 1.20, 1.00, spec["title"],
                 size=54, font=Font.DISPLAY, bold=True, color=Color.WHITE, align="center")
    shapes.label(slide, 0.60, 4.15, G.SLIDE_W - 1.20, 0.40, spec["subtitle"],
                 size=16, font=Font.DISPLAY, italic=True, color=Color.GOLD_LIGHT, align="center")
    shapes.label(
        slide, 0.60, 5.15, G.SLIDE_W - 1.20, 0.36,
        f"**{meta['author']}**~~     NIM {meta['nim']}  ·  {meta['program']}~~",
        size=15, color=Color.WHITE, accent=Color.WHITE, muted=Color.ON_DARK_5, align="center",
    )


LAYOUTS = {"cover": cover, "section": section, "content": content, "closing": closing}
