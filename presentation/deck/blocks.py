"""The block vocabulary.

A slide body is a list of blocks. Each block declares a ``type`` and is
rendered by the matching pair of functions registered in :data:`BLOCKS`:

``measure(spec, width) -> float``   natural height in inches
``render(slide, spec, box)``        draw into an allotted box

Blocks never position themselves vertically — :mod:`deck.layout` does that.
"""

from __future__ import annotations

import math
from typing import Callable, NamedTuple

from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

from . import charts, shapes
from .theme import Color, Font, Grid, Size, Stroke

PAD = Grid.PAD
GUTTER = Grid.GUTTER


class Box(NamedTuple):
    x: float
    y: float
    w: float
    h: float


# --------------------------------------------------------------------------
# Tone helpers — a "tone" is a surface, not a colour
# --------------------------------------------------------------------------
class Tone(NamedTuple):
    fill: str | None
    line: str | None
    title: str
    body: str
    accent: str
    muted: str


TONES = {
    "light": Tone(Color.PANEL, Color.LINE, Color.NAVY, Color.INK, Color.NAVY, Color.MUTED),
    "alt": Tone(Color.PANEL_ALT, Color.LINE, Color.NAVY, Color.INK, Color.NAVY, Color.MUTED),
    "white": Tone(Color.WHITE, Color.LINE, Color.NAVY, Color.INK, Color.NAVY, Color.MUTED),
    "cream": Tone(Color.CREAM, Color.GOLD, Color.NAVY, Color.INK, Color.NAVY, Color.MUTED),
    "dark": Tone(Color.NAVY, Color.NAVY, Color.WHITE, Color.ON_DARK_2, Color.WHITE, Color.ON_DARK_5),
    "accent": Tone(Color.PANEL_ALT, Color.GOLD, Color.NAVY, Color.INK, Color.NAVY, Color.MUTED),
    "plain": Tone(None, None, Color.NAVY, Color.INK, Color.NAVY, Color.MUTED),
}


def tone_of(spec: dict, default: str = "light") -> Tone:
    return TONES[spec.get("tone", default)]


def _tag_color(tag: str) -> str:
    t = tag.strip().upper()
    if t.startswith("EMPIRIS"):
        return Color.GREEN
    if t.startswith("KRITIS") or t.startswith("TERBATAS"):
        return Color.RED
    return Color.GOLD


def _cols(spec: dict, width: float, count: int) -> tuple[list[float], float]:
    """Column x-offsets and the shared column width."""
    gutter = spec.get("gutter", GUTTER)
    col_w = (width - gutter * (count - 1)) / count
    return [i * (col_w + gutter) for i in range(count)], col_w


def _cols_weighted(spec: dict, width: float, count: int) -> tuple[list[float], list[float]]:
    """Unequal columns from a ``widths:`` list of relative weights."""
    gutter = spec.get("gutter", GUTTER)
    weights = list(spec.get("widths") or [1] * count)
    weights = (weights + [1] * count)[:count]
    total = sum(weights) or 1
    usable = width - gutter * (count - 1)
    widths = [usable * w / total for w in weights]
    offsets, cursor = [], 0.0
    for w in widths:
        offsets.append(cursor)
        cursor += w + gutter
    return offsets, widths


# --------------------------------------------------------------------------
# text / lead
# --------------------------------------------------------------------------
def _text_size(spec: dict) -> float:
    return spec.get("size", Size.LEAD if spec["type"] == "lead" else Size.BODY)


def m_text(spec: dict, width: float) -> float:
    size = _text_size(spec)
    font = spec.get("font", Font.BODY)
    body = spec.get("text", spec.get("body", ""))
    paras = [body] if isinstance(body, str) else list(body)
    gap = spec.get("para_gap", 0.10)
    return sum(shapes.text_height(p, width, size, font) for p in paras) + gap * (len(paras) - 1)


def r_text(slide, spec: dict, box: Box) -> None:
    tone = tone_of(spec, spec.get("tone", "plain"))
    body = spec.get("text", spec.get("body", ""))
    shapes.label(
        slide, box.x, box.y, box.w, box.h, body,
        size=_text_size(spec),
        font=spec.get("font", Font.BODY),
        color=spec.get("color", tone.body),
        accent=spec.get("accent", tone.accent),
        muted=tone.muted,
        align=spec.get("align", "left"),
        bold=spec.get("bold", False),
        italic=spec.get("italic", False),
        space_after=spec.get("para_gap", 0.10) * 72,
    )


# --------------------------------------------------------------------------
# heading — a small gold/navy section label inside the body area
# --------------------------------------------------------------------------
def m_heading(spec: dict, width: float) -> float:
    return (spec.get("h", 0.30)
            + (0.26 if spec.get("sub") else 0.0)
            + (0.10 if spec.get("rule") else 0.0))


def r_heading(slide, spec: dict, box: Box) -> None:
    align = spec.get("align", "left")
    shapes.label(
        slide, box.x, box.y, box.w, 0.30, spec["text"],
        size=spec.get("size", 12.5), font=Font.DISPLAY, bold=True,
        color=spec.get("color", Color.NAVY), align=align,
    )
    if spec.get("sub"):
        shapes.label(
            slide, box.x, box.y + 0.30, box.w, 0.24, spec["sub"],
            size=spec.get("sub_size", 10), bold=True,
            color=spec.get("sub_color", Color.GOLD), align=align,
        )
    if spec.get("rule"):
        shapes.hrule(slide, box.x, box.y + box.h - 0.04, box.w,
                     color=spec.get("rule_color", Color.LINE))


# --------------------------------------------------------------------------
# cards — the workhorse: a row of equal panels
# --------------------------------------------------------------------------
def m_cards(spec: dict, width: float) -> float:
    items = spec["items"]
    count = spec.get("columns", len(items))
    _, col_w = _cols(spec, width, count)
    pad = spec.get("pad", PAD)
    inner = col_w - 2 * pad
    badge_h = 0.50 if any(i.get("badge") for i in items) else 0.0

    key_h = 0.40 if any(i.get("key") for i in items) else 0.0

    tallest = 0.0
    for item in items:
        h = badge_h + key_h
        if item.get("title"):
            h += shapes.text_height(item["title"], inner, spec.get("title_size", Size.CARD_TITLE), Font.DISPLAY) + 0.10
        if item.get("body"):
            h += shapes.text_height(item["body"], inner, spec.get("body_size", Size.MICRO))
        tallest = max(tallest, h)
    rows = math.ceil(len(items) / count)
    row_gap = spec.get("row_gap", 0.16)
    return rows * (tallest + 2 * pad) + (rows - 1) * row_gap


def r_cards(slide, spec: dict, box: Box) -> None:
    items = spec["items"]
    count = spec.get("columns", len(items))
    offsets, col_w = _cols(spec, box.w, count)
    pad = spec.get("pad", PAD)
    inner = col_w - 2 * pad
    tone = tone_of(spec)
    badge_h = 0.50 if any(i.get("badge") for i in items) else 0.0
    rows = math.ceil(len(items) / count)
    row_gap = spec.get("row_gap", 0.16)
    card_h = (box.h - (rows - 1) * row_gap) / rows

    for idx, item in enumerate(items):
        col, row = idx % count, idx // count
        x = box.x + offsets[col]
        y = box.y + row * (card_h + row_gap)

        if tone.fill is not None:
            shapes.card(slide, x, y, col_w, card_h, fill=tone.fill, line=tone.line)

        cy = y + pad
        if item.get("badge"):
            bs = 0.44
            shapes.ellipse(slide, x + pad, cy, bs, bs,
                           fill=item.get("badge_fill", Color.NAVY if spec.get("tone") != "dark" else Color.WHITE))
            shapes.label(
                slide, x + pad, cy + 0.06, bs, bs, item["badge"],
                size=13, font=Font.DISPLAY, bold=True, align="center",
                color=item.get("badge_color", Color.GOLD),
            )
        if item.get("tag"):
            tw = 1.60
            shapes.label(
                slide, x + col_w - pad - tw, cy + 0.04, tw, 0.26, item["tag"],
                size=9, font=Font.BODY, bold=True, align="right",
                color=item.get("tag_color", _tag_color(item["tag"])),
            )
        cy += badge_h
        if item.get("key"):
            shapes.label(
                slide, x + pad, cy, inner, 0.34, item["key"],
                size=spec.get("key_size", 15), font=Font.DISPLAY, bold=True,
                color=item.get("key_color", Color.GOLD),
            )
            cy += 0.40

        align = spec.get("align", "left")
        if item.get("title"):
            th = shapes.text_height(item["title"], inner, spec.get("title_size", Size.CARD_TITLE), Font.DISPLAY)
            shapes.label(
                slide, x + pad, cy, inner, th + 0.06, item["title"],
                size=spec.get("title_size", Size.CARD_TITLE), font=Font.DISPLAY, bold=True,
                color=item.get("title_color", tone.title), align=align,
            )
            cy += th + 0.10
        if item.get("body"):
            bh = card_h - (cy - y) - pad
            shapes.label(
                slide, x + pad, cy, inner, max(bh, 0.20), item["body"],
                size=spec.get("body_size", Size.MICRO),
                color=item.get("body_color", tone.muted if spec.get("tone") != "dark" else tone.body),
                accent=tone.accent, muted=tone.muted, align=align,
            )


# --------------------------------------------------------------------------
# stats — big-number cards
# --------------------------------------------------------------------------
def _stat_pad(spec: dict) -> float:
    """A frameless stat row needs no inner padding."""
    return 0.0 if spec.get("tone") == "plain" else spec.get("pad", PAD)


def _paras_height(value, width: float, size: float, font: str = Font.BODY) -> float:
    """Height of a string or a list of strings rendered as paragraphs."""
    paras = [value] if isinstance(value, str) else list(value)
    return sum(shapes.text_height(p, width, size, font) for p in paras)


def _stat_height(spec: dict, item: dict, inner: float) -> float:
    vsize = item.get("value_size", spec.get("value_size", Size.STAT))
    h = vsize * 1.25 / 72 + 0.06
    if item.get("label"):
        h += shapes.text_height(item["label"], inner, 12.0, Font.DISPLAY) + 0.06
    if item.get("note"):
        h += _paras_height(item["note"], inner, spec.get("note_size", Size.TINY))
    return h


def m_stats(spec: dict, width: float) -> float:
    items = spec["items"]
    pad = _stat_pad(spec)
    _, widths = _cols_weighted(spec, width, len(items))
    tallest = max(_stat_height(spec, item, w - 2 * pad) for item, w in zip(items, widths))
    return tallest + 2 * pad


def r_stats(slide, spec: dict, box: Box) -> None:
    items = spec["items"]
    offsets, widths = _cols_weighted(spec, box.w, len(items))
    pad = _stat_pad(spec)
    tone = tone_of(spec)
    dark = spec.get("tone") == "dark"
    align = spec.get("align", "left")

    for idx, item in enumerate(items):
        x = box.x + offsets[idx]
        col_w = widths[idx]
        inner = col_w - 2 * pad
        if tone.fill is not None:
            shapes.card(slide, x, box.y, col_w, box.h, fill=tone.fill, line=tone.line)

        cy = box.y + pad
        if spec.get("valign") == "middle":
            cy = box.y + max(pad, (box.h - _stat_height(spec, item, inner)) / 2)

        vsize = item.get("value_size", spec.get("value_size", Size.STAT))
        vh = vsize * 1.25 / 72
        shapes.label(
            slide, x + pad, cy, inner, vh, item["value"],
            size=vsize, font=Font.DISPLAY, bold=True,
            color=item.get("value_color", Color.GOLD if dark else Color.NAVY),
            align=align, wrap=False,
        )
        cy += vh + 0.06
        if item.get("label"):
            lh = shapes.text_height(item["label"], inner, 12.0, Font.DISPLAY)
            shapes.label(
                slide, x + pad, cy, inner, lh + 0.04, item["label"],
                size=12.0, font=Font.DISPLAY, bold=True,
                color=item.get("label_color", Color.WHITE if dark else Color.GOLD),
                align=align,
            )
            cy += lh + 0.06
        if item.get("note"):
            shapes.label(
                slide, x + pad, cy, inner, box.y + box.h - cy - pad + 0.05, item["note"],
                size=spec.get("note_size", Size.TINY),
                color=item.get("note_color", Color.ON_DARK_4 if dark else Color.MUTED),
                accent=tone.accent, muted=tone.muted, align=align,
            )


# --------------------------------------------------------------------------
# metrics — badge + title + explanation, in two columns
# --------------------------------------------------------------------------
def m_metrics(spec: dict, width: float) -> float:
    items = spec["items"]
    count = spec.get("columns", 2)
    _, col_w = _cols(spec, width, count)
    badge = spec.get("badge_size", 1.18)
    text_w = col_w - badge - 0.16
    row_gap = spec.get("row_gap", 0.22)
    rows = math.ceil(len(items) / count)

    tallest = 0.0
    for item in items:
        h = 0.30 + shapes.text_height(item["body"], text_w, Size.TINY)
        tallest = max(tallest, h)
    tallest = max(tallest, badge)
    return rows * tallest + (rows - 1) * row_gap


def r_metrics(slide, spec: dict, box: Box) -> None:
    items = spec["items"]
    count = spec.get("columns", 2)
    offsets, col_w = _cols(spec, box.w, count)
    badge = spec.get("badge_size", 1.18)
    text_x = badge + 0.16
    text_w = col_w - text_x
    rows = math.ceil(len(items) / count)
    row_gap = spec.get("row_gap", 0.22)
    row_h = (box.h - (rows - 1) * row_gap) / rows
    # column-major so a 2-column guide reads down each side, as in the template
    order = spec.get("order", "column")

    for idx, item in enumerate(items):
        if order == "column":
            col, row = idx // rows, idx % rows
        else:
            col, row = idx % count, idx // count
        x = box.x + offsets[min(col, count - 1)]
        y = box.y + row * (row_h + row_gap)

        bh = min(badge, row_h)
        shapes.card(slide, x, y, badge, bh, fill=Color.NAVY, line=Color.NAVY)
        shapes.label(
            slide, x, y + bh / 2 - 0.20, badge, 0.40, item["value"],
            size=item.get("value_size", Size.STAT_SM), font=Font.DISPLAY, bold=True,
            color=Color.GOLD, align="center",
        )
        head = item["title"] + (f"    **{item['hint']}**" if item.get("hint") else "")
        shapes.label(
            slide, x + text_x, y - 0.03, text_w, 0.30, head,
            size=12.5, font=Font.DISPLAY, bold_font=Font.BODY, bold=False,
            color=Color.NAVY, accent=Color.GOLD,
        )
        shapes.label(
            slide, x + text_x, y + 0.27, text_w, row_h - 0.27, item["body"],
            size=Size.TINY, color=Color.INK,
        )


# --------------------------------------------------------------------------
# banner — a full-width callout strip
# --------------------------------------------------------------------------
def m_banner(spec: dict, width: float) -> float:
    pad = spec.get("pad", PAD)
    label_w = spec.get("label_w", 2.60) if spec.get("label") else 0.0
    inner = width - 2 * pad - (label_w + 0.25 if label_w else 0.0)
    size = spec.get("size", Size.BODY)
    return shapes.text_height(spec["body"], inner, size) + 2 * spec.get("vpad", 0.14)


def r_banner(slide, spec: dict, box: Box) -> None:
    tone = tone_of(spec)
    pad = spec.get("pad", PAD)
    if tone.fill is not None:
        shapes.card(slide, box.x, box.y, box.w, box.h,
                    fill=tone.fill, line=tone.line,
                    line_w=Stroke.THIN)
    tx, tw = box.x + pad + 0.03, box.w - 2 * pad - 0.06
    if spec.get("label"):
        label_w = spec.get("label_w", 2.60)
        shapes.label(
            slide, tx, box.y + spec.get("vpad", 0.14), label_w, box.h - 0.20, spec["label"],
            size=spec.get("label_size", 12.0), font=Font.DISPLAY, bold=True,
            color=spec.get("label_color", Color.GOLD if spec.get("tone") == "dark" else Color.NAVY),
        )
        tx += label_w + 0.25
        tw = box.x + box.w - pad - tx
    shapes.label(
        slide, tx, box.y + spec.get("vpad", 0.14), tw, box.h - 0.20, spec["body"],
        size=spec.get("size", Size.BODY), color=spec.get("color", tone.body),
        accent=tone.accent, muted=tone.muted, align=spec.get("align", "left"),
        anchor=spec.get("anchor", "top"),
    )


# --------------------------------------------------------------------------
# bullets
# --------------------------------------------------------------------------
def m_bullets(spec: dict, width: float) -> float:
    size = spec.get("size", Size.BODY)
    indent = spec.get("indent", 0.24)
    gap = spec.get("gap", 0.14)
    total = 0.0
    for item in spec["items"]:
        total += max(shapes.text_height(item, width - indent, size), size * 1.12 / 72) + gap
    return total - gap


def r_bullets(slide, spec: dict, box: Box) -> None:
    tone = tone_of(spec, "plain")
    size = spec.get("size", Size.BODY)
    indent = spec.get("indent", 0.24)
    gap = spec.get("gap", 0.14)
    marker = spec.get("marker", 0.10)
    y = box.y
    for item in spec["items"]:
        h = max(shapes.text_height(item, box.w - indent, size), size * 1.12 / 72)
        shapes.rect(slide, box.x, y + size / 72 * 0.42, marker, marker,
                    fill=spec.get("marker_color", Color.GOLD))
        shapes.label(
            slide, box.x + indent, y, box.w - indent, h + 0.06, item,
            size=size, color=spec.get("color", tone.body),
            accent=tone.accent, muted=tone.muted,
        )
        y += h + gap


def m_numbered(spec: dict, width: float) -> float:
    size = spec.get("size", 11.3)
    badge = spec.get("badge_size", 0.34)
    indent = badge + 0.12
    gap = spec.get("gap", 0.27)
    total = 0.0
    for item in spec["items"]:
        total += max(shapes.text_height(item, width - indent, size), badge) + gap
    return total - gap


def r_numbered(slide, spec: dict, box: Box) -> None:
    tone = tone_of(spec, "plain")
    size = spec.get("size", 11.3)
    badge = spec.get("badge_size", 0.34)
    indent = badge + 0.12
    gap = spec.get("gap", 0.27)
    y = box.y
    for i, item in enumerate(spec["items"], spec.get("start", 1)):
        h = max(shapes.text_height(item, box.w - indent, size), badge)
        shapes.ellipse(slide, box.x, y, badge, badge, fill=spec.get("badge_fill", Color.GOLD))
        shapes.label(
            slide, box.x, y + 0.055, badge, badge, str(i),
            size=13, font=Font.DISPLAY, bold=True, align="center",
            color=spec.get("badge_color", Color.NAVY),
        )
        shapes.label(
            slide, box.x + indent, y - 0.02, box.w - indent, h + 0.10, item,
            size=size, color=spec.get("color", tone.body), accent=tone.accent, muted=tone.muted,
        )
        y += h + gap


# --------------------------------------------------------------------------
# panel — a titled container holding nested blocks
# --------------------------------------------------------------------------
def m_panel(spec: dict, width: float) -> float:
    from .layout import measure_stack

    pad = spec.get("pad", 0.22)
    head = 0.0
    if spec.get("title"):
        head = 0.30 if not spec.get("band") else 0.50
        head += 0.04
    return 2 * pad + head + measure_stack(spec.get("blocks", []), width - 2 * pad)


def r_panel(slide, spec: dict, box: Box) -> None:
    from .layout import render_stack

    tone = tone_of(spec)
    pad = spec.get("pad", 0.22)
    if tone.fill is not None:
        shapes.card(slide, box.x, box.y, box.w, box.h, fill=tone.fill, line=tone.line)

    y = box.y + pad
    if spec.get("band"):
        band_h = spec.get("band_h", 0.50)
        shapes.card(slide, box.x, box.y, box.w, band_h,
                    fill=spec.get("band_fill", Color.NAVY), line=spec.get("band_fill", Color.NAVY))
        shapes.rect(slide, box.x, box.y + band_h - 0.10, box.w, 0.10,
                    fill=spec.get("band_fill", Color.NAVY))
        shapes.label(
            slide, box.x + pad, box.y + 0.13, box.w - 2 * pad, band_h - 0.16, spec["title"],
            size=spec.get("title_size", Size.PANEL_TITLE), font=Font.DISPLAY, bold=True,
            color=spec.get("title_color", Color.WHITE),
        )
        y = box.y + band_h + 0.18
    elif spec.get("title"):
        shapes.label(
            slide, box.x + pad, y, box.w - 2 * pad, 0.30, spec["title"],
            size=spec.get("title_size", Size.PANEL_TITLE), font=Font.DISPLAY, bold=True,
            color=spec.get("title_color", tone.title if spec.get("tone") != "dark" else Color.GOLD),
        )
        y += 0.34

    render_stack(slide, spec.get("blocks", []),
                 box.x + pad, y, box.w - 2 * pad, box.y + box.h - pad)


# --------------------------------------------------------------------------
# columns — side-by-side regions, each holding nested blocks
# --------------------------------------------------------------------------
def m_columns(spec: dict, width: float) -> float:
    from .layout import measure_stack

    widths = _column_widths(spec, width)
    return max(
        measure_stack(item.get("blocks", []), w)
        for item, w in zip(spec["items"], widths)
    )


def _column_widths(spec: dict, width: float) -> list[float]:
    items = spec["items"]
    gutter = spec.get("gutter", GUTTER)
    declared = [item.get("width") for item in items]
    free = width - gutter * (len(items) - 1) - sum(d for d in declared if d)
    unset = sum(1 for d in declared if not d)
    return [d if d else (free / unset if unset else 0) for d in declared]


def r_columns(slide, spec: dict, box: Box) -> None:
    from .layout import render_stack

    widths = _column_widths(spec, box.w)
    gutter = spec.get("gutter", GUTTER)
    x = box.x
    for item, w in zip(spec["items"], widths):
        render_stack(slide, item.get("blocks", []), x, box.y, w, box.y + box.h)
        x += w + gutter


# --------------------------------------------------------------------------
# table
# --------------------------------------------------------------------------
def m_table(spec: dict, width: float) -> float:
    header = spec.get("header_h", 0.40)
    row_h = spec.get("row_h", 0.40)
    return header + row_h * len(spec["rows"])


def r_table(slide, spec: dict, box: Box) -> None:
    cols = spec["columns"]
    rows = spec["rows"]
    header_h = spec.get("header_h", 0.40)
    row_h = (box.h - header_h) / max(len(rows), 1)

    gfx = slide.shapes.add_table(
        len(rows) + 1, len(cols),
        Inches(box.x), Inches(box.y), Inches(box.w), Inches(box.h),
    )
    table = gfx.table
    table.first_row = False
    table.horz_banding = False

    weights = spec.get("widths") or [1] * len(cols)
    total = sum(weights)
    for i, weight in enumerate(weights):
        table.columns[i].width = Inches(box.w * weight / total)
    table.rows[0].height = Inches(header_h)
    for r in range(1, len(rows) + 1):
        table.rows[r].height = Inches(row_h)

    aligns = spec.get("align") or ["left"] * len(cols)

    def style_cell(cell, text, *, fill, color, bold, size, font, align):
        cell.fill.solid()
        cell.fill.fore_color.rgb = shapes.rgb(fill)
        cell.margin_left = cell.margin_right = Inches(0.083)
        cell.margin_top = cell.margin_bottom = Inches(0.042)
        cell.vertical_anchor = shapes.ANCHOR["middle"]
        shapes.write_text(
            cell.text_frame, text, size=size, color=color, font=font,
            bold=bold, align=align, accent=Color.NAVY, spacing=1.0,
        )
        _cell_borders(cell)

    for c, name in enumerate(cols):
        style_cell(table.cell(0, c), name, fill=Color.NAVY, color=Color.WHITE,
                   bold=True, size=spec.get("header_size", 11.5), font=Font.DISPLAY,
                   align=aligns[c])
    for r, row in enumerate(rows, start=1):
        fill = Color.WHITE if r % 2 else Color.PANEL_ALT
        for c, value in enumerate(row):
            style_cell(table.cell(r, c), str(value), fill=fill, color=Color.INK,
                       bold=False, size=spec.get("size", 10.5), font=Font.BODY,
                       align=aligns[c])


def _cell_borders(cell) -> None:
    from lxml import etree
    from pptx.oxml.ns import qn

    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        for old in tcPr.findall(qn(tag)):
            tcPr.remove(old)
    # order matters in the schema: L, R, T, B come first inside tcPr
    for i, tag in enumerate(("a:lnL", "a:lnR", "a:lnT", "a:lnB")):
        ln = etree.SubElement(tcPr, qn(tag))
        ln.set("w", "6350")
        ln.set("cap", "flat")
        ln.set("cmpd", "sng")
        ln.set("algn", "ctr")
        fill = etree.SubElement(ln, qn("a:solidFill"))
        clr = etree.SubElement(fill, qn("a:srgbClr"))
        clr.set("val", Color.LINE)
        tcPr.insert(i, ln)


# --------------------------------------------------------------------------
# chart
# --------------------------------------------------------------------------
def m_chart(spec: dict, width: float) -> float:
    return spec.get("h", 4.60)


def r_chart(slide, spec: dict, box: Box) -> None:
    if spec.get("frame", True):
        shapes.card(slide, box.x, box.y, box.w, box.h, fill=Color.WHITE, line=Color.LINE)
        inset = spec.get("inset", 0.15)
    else:
        inset = 0.0
    charts.add_bar_chart(
        slide,
        box.x + inset, box.y + inset, box.w - 2 * inset, box.h - 2 * inset,
        categories=spec["categories"],
        values=spec["values"],
        series_name=spec.get("series", "Nilai"),
        number_format=spec.get("format", '0.00"%"'),
        horizontal=spec.get("horizontal", True),
        highlight=spec.get("highlight"),
        labels=spec.get("labels"),
    )


# --------------------------------------------------------------------------
# tiers — the three-tier architecture blueprint
# --------------------------------------------------------------------------
def m_tiers(spec: dict, width: float) -> float:
    return spec.get("h", 3.50)


def r_tiers(slide, spec: dict, box: Box) -> None:
    items = spec["items"]
    arrow_w = spec.get("arrow", 0.34)
    col_w = (box.w - arrow_w * (len(items) - 1)) / len(items)
    band_h = spec.get("band_h", 0.84)
    pad = 0.22

    for idx, item in enumerate(items):
        x = box.x + idx * (col_w + arrow_w)
        shapes.card(slide, x, box.y, col_w, box.h, fill=Color.PANEL, line=Color.LINE)
        shapes.card(slide, x, box.y, col_w, band_h, fill=Color.NAVY, line=Color.NAVY)
        shapes.rect(slide, x, box.y + band_h - 0.12, col_w, 0.12, fill=Color.NAVY)
        shapes.label(slide, x + pad, box.y + 0.12, col_w - 2 * pad, 0.24, item["kicker"],
                     size=11, bold=True, color=Color.GOLD)
        shapes.label(slide, x + pad, box.y + 0.38, col_w - 2 * pad, 0.34, item["title"],
                     size=16, font=Font.DISPLAY, bold=True, color=Color.WHITE)

        y = box.y + band_h + 0.10
        shapes.label(slide, x + pad, y, col_w - 2 * pad, 0.26, item["subtitle"],
                     size=Size.TINY, italic=True, color=Color.MUTED)
        y += 0.36
        for row in item["rows"]:
            shapes.label(slide, x + pad, y, col_w - 2 * pad, 0.22, row["label"],
                         size=9, bold=True, color=Color.GOLD)
            h = shapes.text_height(row["value"], col_w - 2 * pad, Size.BODY)
            shapes.label(slide, x + pad, y + 0.20, col_w - 2 * pad, h + 0.08, row["value"],
                         size=Size.BODY, bold=True, color=Color.INK)
            y += 0.22 + h + 0.16

        if idx < len(items) - 1:
            shapes.label(
                slide, x + col_w + 0.05, box.y + box.h / 2 - 0.24, arrow_w - 0.10, 0.40, "▶",
                size=16, align="center", color=Color.GOLD,
            )


# --------------------------------------------------------------------------
# agenda
# --------------------------------------------------------------------------
def m_agenda(spec: dict, width: float) -> float:
    count = spec.get("columns", 2)
    rows = math.ceil(len(spec["items"]) / count)
    return rows * spec.get("row_h", 1.24) - spec.get("row_gap", 0.0)


def r_agenda(slide, spec: dict, box: Box) -> None:
    items = spec["items"]
    count = spec.get("columns", 2)
    offsets, col_w = _cols(spec, box.w, count)
    rows = math.ceil(len(items) / count)
    row_h = box.h / rows
    badge = 0.50
    text_x = badge + 0.18

    for idx, item in enumerate(items):
        col, row = (idx // rows, idx % rows) if spec.get("order", "column") == "column" \
            else (idx % count, idx // count)
        x = box.x + offsets[min(col, count - 1)]
        y = box.y + row * row_h

        # NB: the field is `num`, not `no` — YAML 1.1 reads a bare `no` as False.
        shapes.ellipse(slide, x, y + 0.06, badge, badge, fill=Color.NAVY)
        shapes.label(slide, x, y + 0.18, badge, badge, str(item["num"]),
                     size=13, font=Font.DISPLAY, bold=True, align="center", color=Color.GOLD)
        shapes.label(slide, x + text_x, y, col_w - text_x, 0.36, item["title"],
                     size=14.5, font=Font.DISPLAY, bold=True, color=Color.NAVY)
        shapes.label(slide, x + text_x, y + 0.36, col_w - text_x, 0.50, item["body"],
                     size=Size.SMALL, color=Color.MUTED)


# --------------------------------------------------------------------------
# glossary — three columns of term / definition pairs
# --------------------------------------------------------------------------
def m_glossary(spec: dict, width: float) -> float:
    rows = max(len(g["items"]) for g in spec["groups"])
    return 0.36 + rows * spec.get("row_h", 0.40)


def r_glossary(slide, spec: dict, box: Box) -> None:
    groups = spec["groups"]
    offsets, col_w = _cols(spec, box.w, len(groups))
    term_w = spec.get("term_w", 1.18)
    row_h = spec.get("row_h", 0.40)

    for gi, group in enumerate(groups):
        x = box.x + offsets[gi]
        shapes.label(slide, x, box.y, col_w, 0.30, group["title"],
                     size=11.5, font=Font.DISPLAY, bold=True, color=Color.GOLD)
        shapes.hrule(slide, x, box.y + 0.33, col_w, color=Color.LINE)
        y = box.y + 0.36
        for term, definition in group["items"]:
            shapes.label(slide, x, y, term_w, row_h, term,
                         size=Size.TINY, bold=True, color=Color.NAVY)
            shapes.label(slide, x + term_w, y, col_w - term_w, row_h, definition,
                         size=9.5, color=Color.INK)
            y += row_h


# --------------------------------------------------------------------------
# references
# --------------------------------------------------------------------------
def m_refs(spec: dict, width: float) -> float:
    rows = max(len(g["items"]) for g in spec["groups"])
    return 0.50 + rows * spec.get("row_h", 0.36)


def r_refs(slide, spec: dict, box: Box) -> None:
    groups = spec["groups"]
    offsets, col_w = _cols(spec, box.w, len(groups))
    for gi, group in enumerate(groups):
        x = box.x + offsets[gi]
        shapes.label(slide, x, box.y, col_w, 0.44, group["title"],
                     size=11, font=Font.DISPLAY, bold=True, color=Color.GOLD)
        shapes.label(
            slide, x, box.y + 0.50, col_w, box.h - 0.50, group["items"],
            size=9.5, color=Color.MUTED, accent=Color.NAVY,
            space_after=spec.get("row_gap", 0.36) * 72 - 9.5 * 1.12,
        )


# --------------------------------------------------------------------------
# figure placeholder
# --------------------------------------------------------------------------
def m_figure(spec: dict, width: float) -> float:
    return spec.get("h", 3.50)


def r_figure(slide, spec: dict, box: Box) -> None:
    if spec.get("image"):
        from PIL import Image as _Image
        from pptx.util import Inches as _In

        # Fit inside the card in both directions, centred — a width-only
        # scale would let a tall image overrun the card and the body band.
        inset = spec.get("inset", 0.14)
        avail_w = box.w - 2 * inset
        avail_h = box.h - 2 * inset
        with _Image.open(spec["image"]) as source:
            aspect = source.height / source.width
        w = avail_w
        h = w * aspect
        if h > avail_h:
            h = avail_h
            w = h / aspect
        x = box.x + (box.w - w) / 2
        y = box.y + (box.h - h) / 2
        # The card hugs the fitted image rather than the whole slot, so an
        # aspect mismatch reads as slide background instead of grey bands.
        shapes.card(slide, x - inset, y - inset, w + 2 * inset, h + 2 * inset,
                    fill=Color.PANEL, line=Color.LINE)
        slide.shapes.add_picture(spec["image"], _In(x), _In(y), width=_In(w))
        return
    shapes.card(slide, box.x, box.y, box.w, box.h, fill=Color.PANEL, line=Color.LINE)
    shapes.label(slide, box.x, box.y + box.h * 0.28, box.w, 0.90, "⊞",
                 size=spec.get("glyph_size", 48), align="center", color=Color.GHOST)
    caption = spec.get("caption")
    if caption:
        lines = [caption] if isinstance(caption, str) else list(caption)
        shapes.label(
            slide, box.x + 0.40, box.y + box.h * 0.28 + 1.00, box.w - 0.80, box.h * 0.35, lines,
            size=12.0, italic=True, color=Color.MUTED, align="center",
        )


# --------------------------------------------------------------------------
# maprows — full-width "key -> claim -> evidence" rows
# --------------------------------------------------------------------------
def m_maprows(spec: dict, width: float) -> float:
    row_h = spec.get("row_h", 1.00)
    gap = spec.get("gap", 0.12)
    return len(spec["items"]) * row_h + (len(spec["items"]) - 1) * gap


def r_maprows(slide, spec: dict, box: Box) -> None:
    items = spec["items"]
    gap = spec.get("gap", 0.12)
    row_h = (box.h - gap * (len(items) - 1)) / len(items)
    key_w = spec.get("key_w", 1.35)
    title_w = spec.get("title_w", 3.60)
    split = spec.get("split", 5.30)

    for i, item in enumerate(items):
        y = box.y + i * (row_h + gap)
        shapes.card(slide, box.x, y, box.w, row_h, fill=Color.PANEL, line=Color.LINE)
        shapes.label(slide, box.x + 0.20, y + 0.18, key_w, row_h - 0.30, item["key"],
                     size=18, font=Font.DISPLAY, bold=True, color=Color.GOLD)
        shapes.label(slide, box.x + key_w + 0.20, y + 0.16, title_w, row_h - 0.28, item["title"],
                     size=13, font=Font.DISPLAY, bold=True, color=Color.NAVY)
        shapes.vrule(slide, box.x + split, y + 0.18, row_h - 0.36, color=Color.LINE)
        shapes.label(slide, box.x + split + 0.25, y + 0.16,
                     box.w - split - 0.45, row_h - 0.28, item["body"],
                     size=Size.BODY, color=Color.INK)


# --------------------------------------------------------------------------
# kvrows — LABEL | value rows used inside the methodology panels
# --------------------------------------------------------------------------
def m_kvrows(spec: dict, width: float) -> float:
    label_w = spec.get("label_w", 1.15)
    size = spec.get("size", Size.SMALL)
    gap = spec.get("gap", 0.16)
    total = 0.0
    for item in spec["items"]:
        total += max(shapes.text_height(item["value"], width - label_w - 0.10, size), 0.26) + gap
    return total - gap


def r_kvrows(slide, spec: dict, box: Box) -> None:
    label_w = spec.get("label_w", 1.15)
    size = spec.get("size", Size.SMALL)
    gap = spec.get("gap", 0.16)
    tone = tone_of(spec, "plain")
    y = box.y
    for item in spec["items"]:
        h = max(shapes.text_height(item["value"], box.w - label_w - 0.10, size), 0.26)
        shapes.label(slide, box.x, y + 0.01, label_w, h, item["label"],
                     size=10, bold=True, color=Color.GOLD)
        shapes.label(slide, box.x + label_w + 0.10, y, box.w - label_w - 0.10, h + 0.06,
                     item["value"], size=size, color=spec.get("color", tone.body),
                     accent=tone.accent, muted=tone.muted)
        y += h + gap


# --------------------------------------------------------------------------
# source note & spacer
# --------------------------------------------------------------------------
def m_source(spec: dict, width: float) -> float:
    return 0.26


def r_source(slide, spec: dict, box: Box) -> None:
    text = f"**{spec.get('label', 'Sumber dan metode:')}**  ~~{spec['body']}~~"
    shapes.label(slide, box.x, box.y, box.w, 0.26, text,
                 size=Size.SOURCE, color=Color.MUTED, accent=Color.MUTED, italic=False)


def m_spacer(spec: dict, width: float) -> float:
    return spec.get("h", 0.20)


def r_spacer(slide, spec: dict, box: Box) -> None:
    return None


def m_rule(spec: dict, width: float) -> float:
    return spec.get("h", 0.02)


def r_rule(slide, spec: dict, box: Box) -> None:
    shapes.hrule(slide, box.x, box.y, box.w, color=spec.get("color", Color.LINE))


# --------------------------------------------------------------------------
# Registry
# --------------------------------------------------------------------------
BLOCKS: dict[str, tuple[Callable, Callable]] = {
    "lead": (m_text, r_text),
    "text": (m_text, r_text),
    "heading": (m_heading, r_heading),
    "cards": (m_cards, r_cards),
    "stats": (m_stats, r_stats),
    "metrics": (m_metrics, r_metrics),
    "banner": (m_banner, r_banner),
    "bullets": (m_bullets, r_bullets),
    "numbered": (m_numbered, r_numbered),
    "panel": (m_panel, r_panel),
    "columns": (m_columns, r_columns),
    "table": (m_table, r_table),
    "chart": (m_chart, r_chart),
    "tiers": (m_tiers, r_tiers),
    "agenda": (m_agenda, r_agenda),
    "glossary": (m_glossary, r_glossary),
    "refs": (m_refs, r_refs),
    "figure": (m_figure, r_figure),
    "maprows": (m_maprows, r_maprows),
    "kvrows": (m_kvrows, r_kvrows),
    "source": (m_source, r_source),
    "spacer": (m_spacer, r_spacer),
    "rule": (m_rule, r_rule),
}
