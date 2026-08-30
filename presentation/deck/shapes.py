"""Low-level drawing primitives.

Everything above this layer speaks in inches and design tokens; this module
is the only place that touches raw OXML.
"""

from __future__ import annotations

import math
import re
from typing import NamedTuple

from lxml import etree
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from .theme import LINE_SPACING, Color, Font, Grid, Size, Stroke, rgb

ALIGN = {
    "left": PP_ALIGN.LEFT,
    "center": PP_ALIGN.CENTER,
    "right": PP_ALIGN.RIGHT,
    "justify": PP_ALIGN.JUSTIFY,
}

ANCHOR = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}


# --------------------------------------------------------------------------
# Inline markup
# --------------------------------------------------------------------------
class Run(NamedTuple):
    text: str
    bold: bool = False
    italic: bool = False
    color: str | None = None
    size: float | None = None
    font: str | None = None


_TOKEN = re.compile(r"(\*\*\*|\*\*|\*|~~)")


def parse_markup(text: str) -> list[Run]:
    """Turn a lightweight markup string into styled runs.

    ``**bold**`` · ``*italic*`` · ``***bold italic***`` · ``~~muted~~``

    Bold is the emphasis carrier in this template: bold runs also take the
    accent colour, which is what gives the reference deck its two-tone
    paragraphs. Colour is resolved later by :func:`write_text`, which knows
    whether the surface is light or dark.
    """
    runs: list[Run] = []
    bold = italic = muted = False
    buf: list[str] = []

    def flush() -> None:
        if buf:
            runs.append(Run("".join(buf), bold, italic, "~muted" if muted else None))
            buf.clear()

    for part in _TOKEN.split(text):
        if part == "***":
            flush()
            bold = not bold
            italic = not italic
        elif part == "**":
            flush()
            bold = not bold
        elif part == "*":
            flush()
            italic = not italic
        elif part == "~~":
            flush()
            muted = not muted
        elif part:
            buf.append(part)
    flush()
    return runs or [Run("")]


def plain(text: str) -> str:
    """Markup stripped — used for width measurement."""
    return _TOKEN.sub("", text)


# --------------------------------------------------------------------------
# Measurement
# --------------------------------------------------------------------------
def chars_per_line(width_in: float, size_pt: float, font: str = Font.BODY) -> float:
    advance = Font.ADVANCE.get(font, 0.48) * size_pt
    return max(1.0, width_in * 72.0 / advance)


def estimate_lines(text: str, width_in: float, size_pt: float, font: str = Font.BODY) -> int:
    """Line count for `text` wrapped into `width_in`, honouring explicit \n."""
    cpl = chars_per_line(width_in, size_pt, font)
    total = 0
    for para in plain(text).split("\n"):
        total += max(1, math.ceil(len(para) / cpl))
    return total


def text_height(
    text: str,
    width_in: float,
    size_pt: float,
    font: str = Font.BODY,
    spacing: float = LINE_SPACING,
) -> float:
    """Estimated rendered height in inches."""
    return estimate_lines(text, width_in, size_pt, font) * size_pt * spacing / 72.0


# --------------------------------------------------------------------------
# OXML helpers
# --------------------------------------------------------------------------
def _spPr(shape):
    return shape._element.spPr


def _no_shadow(shape) -> None:
    """Suppress the theme's inherited shadow; the template is flat."""
    spPr = _spPr(shape)
    for existing in spPr.findall(qn("a:effectLst")):
        spPr.remove(existing)
    etree.SubElement(spPr, qn("a:effectLst"))


def _round_adj(shape, w_in: float, h_in: float, radius_in: float) -> None:
    """Set a roundRect's corner radius to an absolute length."""
    prstGeom = _spPr(shape).find(qn("a:prstGeom"))
    if prstGeom is None:
        return
    for gd in prstGeom.findall(".//" + qn("a:gd")):
        gd.getparent().remove(gd)
    avLst = prstGeom.find(qn("a:avLst"))
    if avLst is None:
        avLst = etree.SubElement(prstGeom, qn("a:avLst"))
    gd = etree.SubElement(avLst, qn("a:gd"))
    gd.set("name", "adj")
    adj = int(round(radius_in / max(min(w_in, h_in), 1e-6) * 100000))
    gd.set("fmla", f"val {max(0, min(adj, 50000))}")


def set_background(slide, color: str) -> None:
    bg = slide._element.find(qn("p:bg"))
    if bg is not None:
        slide._element.remove(bg)
    bg = etree.Element(qn("p:bg"))
    bgPr = etree.SubElement(bg, qn("p:bgPr"))
    fill = etree.SubElement(bgPr, qn("a:solidFill"))
    clr = etree.SubElement(fill, qn("a:srgbClr"))
    clr.set("val", color)
    etree.SubElement(bgPr, qn("a:effectLst"))
    slide._element.find(qn("p:cSld")).insert(0, bg)


# --------------------------------------------------------------------------
# Shape factories
# --------------------------------------------------------------------------
def add_shape(
    slide,
    kind,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    fill: str | None = None,
    line: str | None = None,
    line_w=Stroke.THIN,
    radius: float | None = None,
):
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.shadow.inherit = False
    _no_shadow(shape)

    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(fill)

    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = rgb(line)
        shape.line.width = line_w

    if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        _round_adj(shape, w, h, Grid.RADIUS if radius is None else radius)

    shape.text_frame.word_wrap = True
    return shape


def card(slide, x, y, w, h, *, fill=Color.PANEL, line=Color.LINE, line_w=Stroke.THIN, radius=None):
    return add_shape(
        slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h,
        fill=fill, line=line, line_w=line_w, radius=radius,
    )


def rect(slide, x, y, w, h, *, fill=None, line=None, line_w=Stroke.THIN):
    return add_shape(slide, MSO_SHAPE.RECTANGLE, x, y, w, h, fill=fill, line=line, line_w=line_w)


def ellipse(slide, x, y, w, h, *, fill=Color.NAVY, line=None, line_w=Stroke.THIN):
    return add_shape(slide, MSO_SHAPE.OVAL, x, y, w, h, fill=fill, line=line, line_w=line_w)


def hrule(slide, x, y, w, *, color=Color.LINE, width=Stroke.THIN):
    shape = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(y), Inches(x + w), Inches(y)
    )
    shape.line.color.rgb = rgb(color)
    shape.line.width = width
    return shape


def vrule(slide, x, y, h, *, color=Color.LINE, width=Stroke.THIN):
    shape = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(y), Inches(x), Inches(y + h)
    )
    shape.line.color.rgb = rgb(color)
    shape.line.width = width
    return shape


def textbox(slide, x, y, w, h):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return box


# --------------------------------------------------------------------------
# Text writing
# --------------------------------------------------------------------------
def write_text(
    target,
    content,
    *,
    size: float = Size.BODY,
    color: str = Color.INK,
    accent: str | None = None,
    muted: str | None = None,
    font: str = Font.BODY,
    bold_font: str | None = None,
    bold: bool = False,
    italic: bool = False,
    align: str = "left",
    anchor: str = "top",
    spacing: float = LINE_SPACING,
    space_after: float = 0.0,
    clear: bool = True,
    wrap: bool = True,
):
    """Write one or more paragraphs of marked-up text into a shape.

    `content` may be a string (one paragraph), or a list of strings
    (consecutive paragraphs). `accent` is the colour applied to **bold**
    runs; `muted` the colour for ``~~…~~`` runs. Both default sensibly for
    the given base `color`.
    """
    tf = target.text_frame if hasattr(target, "text_frame") else target
    tf.word_wrap = wrap
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = ANCHOR[anchor]

    if accent is None:
        accent = Color.WHITE if _is_dark(color) else Color.NAVY
    if muted is None:
        muted = Color.ON_DARK_5 if _is_dark(color) else Color.MUTED

    paragraphs = [content] if isinstance(content, str) else list(content)

    if clear:
        tf.clear()

    for i, raw in enumerate(paragraphs):
        para = tf.paragraphs[0] if (i == 0 and clear) else tf.add_paragraph()
        para.alignment = ALIGN[align]
        para.line_spacing = spacing
        if space_after:
            para.space_after = Pt(space_after)

        for run_spec in parse_markup(raw if isinstance(raw, str) else str(raw)):
            run = para.add_run()
            run.text = run_spec.text
            f = run.font
            # Glyphs such as ↑ ↓ ≈ are missing from Cambria; `bold_font` lets a
            # Cambria heading carry its Calibri annotation in the same box.
            f.name = bold_font if (bold_font and run_spec.bold) else font
            f.size = Pt(size if run_spec.size is None else run_spec.size)
            f.bold = bold or run_spec.bold
            f.italic = italic or run_spec.italic
            if run_spec.color == "~muted":
                f.color.rgb = rgb(muted)
            elif run_spec.bold and not bold:
                f.color.rgb = rgb(accent)
            else:
                f.color.rgb = rgb(color)
    return tf


def label(
    slide,
    x,
    y,
    w,
    h,
    content,
    **kwargs,
):
    """Shorthand: a text box plus :func:`write_text`."""
    box = textbox(slide, x, y, w, h)
    write_text(box, content, **kwargs)
    return box


def _is_dark(hex_str: str) -> bool:
    """Perceptual luminance test — decides default accent colours."""
    r, g, b = (int(hex_str[i : i + 2], 16) for i in (0, 2, 4))
    return (0.2126 * r + 0.7152 * g + 0.0722 * b) > 140  # text is light => surface is dark
