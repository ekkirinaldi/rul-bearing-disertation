#!/usr/bin/env python3
"""Runner for the dissertation presentation service.

    ./build.py build     [--spec content/sidang-terbuka.yaml] [--out out/…pptx]
    ./build.py lint      [--spec …]          structural + overflow checks
    ./build.py preview   [--spec …] [--png]  render via LibreOffice
    ./build.py inspect   <file.pptx>         dump an existing deck's design
"""

from __future__ import annotations

import argparse
import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT))

from deck import build_deck, lint_spec, load_spec  # noqa: E402

DEFAULT_SPEC = ROOT / "content" / "sidang-terbuka.yaml"
DEFAULT_OUT = ROOT / "out" / "Sidang_Disertasi_Toto_Suharto.pptx"

SOFFICE_CANDIDATES = [
    "soffice",
    "/Applications/LibreOffice.app/Contents/MacOS/soffice",
    "/usr/bin/soffice",
]


def _soffice() -> str | None:
    for candidate in SOFFICE_CANDIDATES:
        found = shutil.which(candidate) or (candidate if Path(candidate).exists() else None)
        if found:
            return found
    return None


def cmd_lint(args) -> int:
    spec = load_spec(args.spec)
    problems = lint_spec(spec)
    slides = len(spec.get("slides", []))
    if problems:
        for problem in problems:
            print(f"  [WARN] {problem}")
        print(f"\n{len(problems)} problem(s) across {slides} slide(s).")
        return 1
    print(f"OK — {slides} slides, no layout problems.")
    return 0


def cmd_build(args) -> int:
    spec = load_spec(args.spec)
    problems = lint_spec(spec)
    for problem in problems:
        print(f"  [WARN] {problem}")
    if problems and args.strict:
        print("\nRefusing to build with --strict.")
        return 1
    out = build_deck(spec, args.out)
    size_kb = out.stat().st_size / 1024
    print(f"Wrote {out}  ({len(spec.get('slides', []))} slides, {size_kb:.0f} KB)")
    return 0


def cmd_preview(args) -> int:
    rc = cmd_build(args)
    if rc and args.strict:
        return rc
    soffice = _soffice()
    if not soffice:
        print("LibreOffice not found — install it to render previews.")
        return 1
    out_dir = Path(args.out).parent / "preview"
    out_dir.mkdir(parents=True, exist_ok=True)
    fmt = "png" if args.png else "pdf"
    subprocess.run(
        [soffice, "--headless", "--convert-to", fmt, "--outdir", str(out_dir), str(args.out)],
        check=True,
        stdout=subprocess.DEVNULL,
    )
    print(f"Preview → {out_dir}")
    return 0


def cmd_inspect(args) -> int:
    """Dump geometry, colours and type of an existing deck.

    This is how the template was reverse-engineered; keep it around so the
    design tokens can be re-derived if the template ever changes.
    """
    from pptx import Presentation

    prs = Presentation(args.file)
    print(f"{prs.slide_width / 914400:.2f} x {prs.slide_height / 914400:.2f} in, "
          f"{len(prs.slides.__iter__.__self__._sldIdLst)} slides")
    for i, slide in enumerate(prs.slides, 1):
        print(f"\n--- slide {i} ---")
        for shape in slide.shapes:
            geo = f"({shape.left / 914400:.2f},{shape.top / 914400:.2f}) " \
                  f"{shape.width / 914400:.2f}x{shape.height / 914400:.2f}"
            print(f"  {shape.shape_type} {geo}")
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    for run in para.runs:
                        f = run.font
                        colour = ""
                        try:
                            colour = f"#{f.color.rgb}"
                        except Exception:
                            pass
                        print(f"      {run.text!r} {f.name} {f.size and f.size.pt}pt "
                              f"b={f.bold} i={f.italic} {colour}")
    return 0


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__,
                                     formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = parser.add_subparsers(dest="cmd", required=True)

    def add_common(p):
        p.add_argument("--spec", default=str(DEFAULT_SPEC), help="deck YAML")
        p.add_argument("--out", default=str(DEFAULT_OUT), help="output .pptx")
        p.add_argument("--strict", action="store_true", help="fail on layout warnings")

    p_build = sub.add_parser("build", help="render the deck")
    add_common(p_build)
    p_build.set_defaults(func=cmd_build)

    p_lint = sub.add_parser("lint", help="check the spec without rendering")
    p_lint.add_argument("--spec", default=str(DEFAULT_SPEC))
    p_lint.set_defaults(func=cmd_lint)

    p_prev = sub.add_parser("preview", help="build then convert with LibreOffice")
    add_common(p_prev)
    p_prev.add_argument("--png", action="store_true", help="PNG instead of PDF")
    p_prev.set_defaults(func=cmd_preview)

    p_insp = sub.add_parser("inspect", help="dump an existing .pptx")
    p_insp.add_argument("file")
    p_insp.set_defaults(func=cmd_inspect)

    args = parser.parse_args()
    return args.func(args)


if __name__ == "__main__":
    raise SystemExit(main())
