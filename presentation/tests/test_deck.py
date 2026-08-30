"""Regression tests for the presentation builder.

Run with `make test` from presentation/.
"""

from __future__ import annotations

import sys
from pathlib import Path

import pytest
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from deck import build_deck, lint_spec, load_spec  # noqa: E402
from deck.blocks import BLOCKS  # noqa: E402
from deck.shapes import parse_markup, plain, text_height  # noqa: E402
from deck.theme import Color, Grid  # noqa: E402

SPEC_PATH = ROOT / "content" / "sidang-terbuka.yaml"


@pytest.fixture(scope="module")
def spec() -> dict:
    return load_spec(SPEC_PATH)


@pytest.fixture(scope="module")
def deck(spec, tmp_path_factory) -> Presentation:
    out = tmp_path_factory.mktemp("deck") / "deck.pptx"
    build_deck(spec, out)
    return Presentation(str(out))


# --------------------------------------------------------------------------
# Markup
# --------------------------------------------------------------------------
def test_markup_splits_emphasis():
    runs = parse_markup("plain **bold** *italic* ***both***")
    assert [r.text for r in runs] == ["plain ", "bold", " ", "italic", " ", "both"]
    assert [r.bold for r in runs] == [False, True, False, False, False, True]
    assert [r.italic for r in runs] == [False, False, False, True, False, True]


def test_markup_muted_run():
    runs = parse_markup("kata ~~redup~~")
    assert runs[1].color == "~muted"


def test_plain_strips_markers():
    assert plain("**a** *b* ~~c~~") == "a b c"


def test_text_height_grows_with_content():
    short = text_height("satu baris", 6.0, 11.5)
    long = text_height("kata " * 200, 6.0, 11.5)
    assert long > short > 0


def test_explicit_newlines_count_as_lines():
    assert text_height("a\nb\nc", 6.0, 12.0) > text_height("a", 6.0, 12.0)


# --------------------------------------------------------------------------
# Spec integrity
# --------------------------------------------------------------------------
def test_spec_lints_clean(spec):
    assert lint_spec(spec) == []


def test_every_block_type_is_registered(spec):
    def walk(blocks):
        for block in blocks:
            if not isinstance(block, dict):
                continue
            if "type" in block:
                yield block["type"]
            for child in (block.get("blocks") or []):
                yield from walk([child])
            for item in (block.get("items") or []):
                if isinstance(item, dict) and item.get("blocks"):
                    yield from walk(item["blocks"])

    for slide in spec["slides"]:
        for kind in walk(slide.get("blocks", [])):
            assert kind in BLOCKS, kind


def test_no_ampersand_conjunction(spec):
    """ITB Pedoman: `&` may never stand in for `dan` (see CLAUDE.md)."""
    import yaml

    raw = yaml.safe_dump(spec, allow_unicode=True)
    assert " & " not in raw


def test_decimal_comma_in_numbers(spec):
    """Indonesian decimal separator is a comma, never a period."""
    import re
    import yaml

    raw = yaml.safe_dump(spec, allow_unicode=True)
    # a digit.digit pair that is not part of a version, ratio or layout value
    offenders = [
        m.group(0)
        for m in re.finditer(r"\d+\.\d+\s*(?:%|pp\b)", raw)
    ]
    assert offenders == [], offenders


# --------------------------------------------------------------------------
# Rendered output
# --------------------------------------------------------------------------
def test_slide_count_matches_spec(spec, deck):
    assert len(deck.slides._sldIdLst) == len(spec["slides"])


def test_canvas_is_16_by_9(deck):
    assert round(deck.slide_width / 914400, 2) == 13.33
    assert round(deck.slide_height / 914400, 2) == 7.5


def test_nothing_escapes_the_canvas(deck):
    """No shape may start off-slide or overrun the right/bottom edge."""
    tol = 0.02
    for index, slide in enumerate(deck.slides, 1):
        for shape in slide.shapes:
            if shape.left is None or shape.width is None:
                continue
            left = shape.left / 914400
            top = shape.top / 914400
            right = left + shape.width / 914400
            bottom = top + shape.height / 914400
            assert left >= -tol, f"slide {index}: {shape.name} left={left:.2f}"
            assert top >= -tol, f"slide {index}: {shape.name} top={top:.2f}"
            assert right <= Grid.SLIDE_W + tol, f"slide {index}: {shape.name} right={right:.2f}"
            assert bottom <= Grid.SLIDE_H + tol, f"slide {index}: {shape.name} bottom={bottom:.2f}"


def test_content_slides_carry_footer_and_page_number(spec, deck):
    for index, (slide_spec, slide) in enumerate(zip(spec["slides"], deck.slides), 1):
        if slide_spec.get("layout", "content") in ("cover", "closing"):
            continue
        texts = [s.text_frame.text for s in slide.shapes if s.has_text_frame]
        assert any("NIM 33420002" in t for t in texts), f"slide {index} missing footer"
        assert str(index) in texts, f"slide {index} missing page number"


def test_palette_is_closed(deck):
    """Every explicit colour must come from the theme."""
    allowed = {
        getattr(Color, name) for name in dir(Color)
        if name.isupper() and isinstance(getattr(Color, name), str)
    }
    for index, slide in enumerate(deck.slides, 1):
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    try:
                        value = str(run.font.color.rgb)
                    except (AttributeError, TypeError):
                        continue
                    assert value in allowed, f"slide {index}: stray colour #{value}"


def test_chart_labels_use_decimal_comma(deck):
    for slide in deck.slides:
        for shape in slide.shapes:
            if not shape.has_chart:
                continue
            xml = shape.chart._chartSpace.xml
            assert "99,87%" in xml, "literal Indonesian labels missing from chart"


def test_tables_have_a_navy_header_row(deck):
    seen = 0
    for slide in deck.slides:
        for shape in slide.shapes:
            if not shape.has_table:
                continue
            seen += 1
            header = shape.table.cell(0, 0)
            assert str(header.fill.fore_color.rgb) == Color.NAVY
    assert seen >= 2, "expected at least two tables in the deck"
